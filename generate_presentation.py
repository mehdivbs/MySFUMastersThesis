from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.dml.color import RGBColor

ROOT = Path(r"D:\Academics\SFU Thesis\MySFUMastersThesis")
OUT = ROOT / "Human-Centric_Data-Driven_Decision_Support_20Slides.pptx"
FIG = ROOT / "Figures"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLUE = RGBColor(0, 91, 130)
NAVY = RGBColor(13, 31, 62)
TEAL = RGBColor(0, 156, 148)
GRAY = RGBColor(88, 96, 108)
LIGHT = RGBColor(245, 247, 249)
WHITE = RGBColor(255, 255, 255)


def add_slide(title, subtitle="", footer=True):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.25), Inches(0.55))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = 2
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = 180000
        run.font.bold = True
        run.font.color.rgb = NAVY

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(12.25), Inches(0.35))
        p = sub.text_frame.paragraphs[0]
        p.text = subtitle
        p.alignment = 2
        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = 80000
            run.font.color.rgb = GRAY

    if footer:
        foot = slide.shapes.add_textbox(Inches(0.5), Inches(7.1), Inches(12.25), Inches(0.2))
        p = foot.text_frame.paragraphs[0]
        p.text = "Human-Centric Data-Driven Decision Support | SFU Master’s Thesis"
        p.alignment = 2
        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = 70000
            run.font.color.rgb = GRAY
    return slide


def add_bullets(slide, x, y, w, h, title, items):
    box = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.fill.background()
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = 2
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = 180000
        run.font.bold = True
        run.font.color.rgb = BLUE
    for item in items:
        p = tf.add_paragraph()
        p.text = "• " + item
        p.alignment = 1
        p.level = 0
        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = 125000
            run.font.color.rgb = NAVY


def add_text(slide, x, y, w, h, text, size=18, color=NAVY, bold=False, align=1):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = size * 1000
        run.font.bold = bold
        run.font.color.rgb = color


def add_image(slide, name, x, y, w, h):
    slide.shapes.add_picture(str(FIG / name), Inches(x), Inches(y), Inches(w), Inches(h))


def add_chart(slide, x, y, w, h, labels, values, title):
    chart = slide.shapes.add_chart(Inches(x), Inches(y), Inches(w), Inches(h), chart_type=4).chart
    chart.has_title = True
    chart.chart_title.text = title
    chart.series[0].name = "Importance"
    chart.series[0].values = values
    chart.series[0].categories = labels
    chart.series[0].format.fill.fore_color.rgb = TEAL
    chart.series[0].format.line.color.rgb = TEAL

slides = [
    ("Human-Centric Design for Data-Driven Decision Support", "A 20-slide presentation of the SFU Master’s thesis", False),
    ("Thesis at a Glance", "From fragmented tools to an integrated decision-support framework", True),
    ("The Research Problem", "Decision support is fragmented and hard to trust", True),
    ("Research Questions", "Four questions guide the thesis", True),
    ("Research Methodology", "Literature review, thematic analysis, and synthesis", True),
    ("The Literature Landscape", "Six major classes of data-driven techniques", True),
    ("The Decision-Making Process", "Simon’s three-stage model", True),
    ("Machine Learning", "Prediction, classification, grouping, and adaptation", True),
    ("Traditional Statistics", "Interpretability, hypotheses, forecasting, and risk", True),
    ("Data Visualization", "Interactive visual analytics for sense-making", True),
    ("Optimization and Simulation", "Trade-offs, scenarios, and uncertainty", True),
    ("Causal Reasoning", "From correlation to intervention and explanation", True),
    ("Counterfactual Thinking", "Comparing real outcomes with hypothetical alternatives", True),
    ("Humanistic Factors", "Cognition, knowledge, expertise, and decision style", True),
    ("Human-and-Task-Centered Framework", "Aligning tools with the decision stage and user", True),
    ("Intelligence Stage", "Detecting signals and framing the problem", True),
    ("Design Stage", "Defining criteria, alternatives, and causal structures", True),
    ("Choice Stage", "Evaluating outcomes, trade-offs, and sensitivity", True),
    ("ICADOS Architecture", "A conceptual system integrating data-driven tools and causal reasoning", True),
    ("Key Takeaways", "A practical agenda for future decision-support design", False),
]

slide = add_slide(*slides[0])
add_text(slide, 0.6, 1.6, 11.9, 0.5, "Mehdi Vahabisani", 30, bold=True, align=2)
add_text(slide, 0.6, 2.2, 11.9, 0.3, "Master of Science Thesis", 18, color=GRAY, align=2)
tag = slide.shapes.add_shape(1, Inches(1.7), Inches(3.3), Inches(9.8), Inches(0.7))
tag.fill.solid(); tag.fill.fore_color.rgb = TEAL; tag.line.fill.background()
add_text(slide, 1.7, 3.45, 9.8, 0.4, "Data + human judgment + causal understanding", 18, color=WHITE, bold=True, align=1)
add_text(slide, 0.6, 5.1, 11.9, 0.35, "Decision intelligence • Causal reasoning • Human-centered design", 14, color=BLUE, bold=True, align=2)

for index, (title, subtitle, footer) in enumerate(slides[1:], start=1):
    slide = add_slide(title, subtitle, footer)
    if index == 1:
        add_bullets(slide, 0.5, 1.1, 4.4, 4.5, "Core argument", ["Decision support must be task-aware", "No single technique is sufficient", "Human agency remains essential", "Causal, predictive, and optimization tools complement each other"])
        add_bullets(slide, 5.3, 1.1, 6.2, 4.5, "Research contribution", ["The framework integrates methods across stages", "The thesis is conceptual and prescriptive", "The system should support reasoning, not replace it", "The design is intended for organizational decisions"])
    elif index == 2:
        add_bullets(slide, 0.5, 1.1, 4.7, 4.6, "Core challenges", ["Data tools are applied separately", "Correlations can be mistaken for causes", "High-dimensional outputs increase cognitive load", "Black-box models may reduce trust"])
        add_image(slide, "DDDMProcess.png", 5.7, 1.2, 5.6, 3.8)
    elif index == 3:
        add_bullets(slide, 0.5, 1.1, 4.7, 4.6, "Four research questions", ["RQ1: Which tools support decisions?", "RQ2: How do they augment decision-making?", "RQ3: What role does causal reasoning play?", "RQ4: How should tools be integrated for the user and task?"])
        add_image(slide, "ResearchProcess.png", 5.7, 1.2, 5.6, 3.8)
    elif index == 4:
        add_bullets(slide, 0.5, 1.1, 4.7, 4.6, "Methodological approach", ["Review data-driven decision-making literature", "Classify recurring techniques", "Analyze causal reasoning and human factors", "Synthesize a human-and-task-centered framework"])
        add_image(slide, "ResearchProcess.png", 5.7, 1.2, 5.6, 3.8)
    elif index == 5:
        add_bullets(slide, 0.5, 1.1, 2.7, 4.2, "Technique classes", ["Machine learning", "Traditional statistics", "Data visualization", "Natural language models", "Mathematical optimization", "Simulation"])
        add_image(slide, "DDDMLitReview_search_summary.png", 3.5, 1.2, 7.6, 3.7)
    elif index == 6:
        add_bullets(slide, 0.5, 1.1, 3.3, 4.0, "Three stages", ["Intelligence: detect and frame the problem", "Design: define criteria, variables, and alternatives", "Choice: evaluate, compare, and select"])
        add_image(slide, "DDTAssignment.png", 4.2, 1.2, 7.0, 3.8)
    elif index == 7:
        add_bullets(slide, 0.5, 1.1, 3.5, 4.2, "Machine learning", ["Supervised learning: prediction & classification", "Unsupervised learning: clustering & anomalies", "Reinforcement learning: adaptive policies", "Explainable models: transparent recommendations"])
        add_image(slide, "MCDMTree.jpg", 4.35, 1.2, 6.8, 3.6)
    elif index == 8:
        add_bullets(slide, 0.5, 1.1, 3.6, 4.2, "Traditional statistics", ["Descriptive exploration", "Hypothesis testing", "Time-series forecasting", "Risk estimation and threshold monitoring"])
        add_image(slide, "ContingencyTables.png", 4.5, 1.15, 6.2, 3.5)
    elif index == 9:
        add_bullets(slide, 0.5, 1.1, 3.6, 4.2, "Data visualization", ["Pattern recognition", "Alternative comparison", "Temporal and spatial exploration", "Interactive cross-filtering"])
        add_image(slide, "VADashboardPublicHealth.png", 4.5, 1.15, 6.2, 3.5)
    elif index == 10:
        add_bullets(slide, 0.5, 1.1, 3.5, 4.2, "Optimization and simulation", ["Multi-objective optimization: Pareto trade-offs", "MCDM: ranking across criteria", "Simulation: scenario analysis", "Stochastic models: represent uncertainty"])
        add_image(slide, "Biobjective_VIZ.png", 4.4, 1.15, 6.5, 3.7)
    elif index == 11:
        add_bullets(slide, 0.5, 1.1, 3.5, 4.2, "Causal reasoning", ["Causes, effects, mediators, and confounders", "Action-oriented reasoning", "Counterfactual what-if thinking", "Causal models support interventions"])
        add_image(slide, "CausalGraphConstructionProcess.png", 4.3, 1.2, 6.6, 3.6)
    elif index == 12:
        add_bullets(slide, 0.5, 1.1, 3.5, 4.2, "Counterfactual questions", ["What would have happened without this action?", "Which action would have produced a better outcome?", "How robust is the selected option?", "How does uncertainty affect the comparison?"])
        add_image(slide, "DAGExample.png", 4.3, 1.2, 6.6, 3.7)
    elif index == 13:
        add_bullets(slide, 0.5, 1.1, 3.6, 4.2, "Humanistic factors", ["Working memory and attention", "Domain knowledge and mental models", "Expertise and experience", "Decision style and cognitive preferences"])
        add_bullets(slide, 4.5, 1.2, 6.2, 3.9, "Design implications", ["Analytical users benefit from structured evidence", "Intuitive users prefer fast guidance", "Adaptive experts revise assumptions", "Support should not enforce one style"])
    elif index == 14:
        add_image(slide, "ThreeLayerIntegration.png", 0.5, 1.1, 12.0, 4.1)
    elif index == 15:
        add_image(slide, "IntelligenceWorkflow.png", 0.4, 1.1, 12.2, 4.1)
    elif index == 16:
        add_image(slide, "DesignWorkflow_ArrowsAmplified.png", 0.4, 1.1, 12.2, 4.1)
    elif index == 17:
        add_image(slide, "ChoiceWorkflow.png", 0.4, 1.1, 12.2, 4.1)
    elif index == 18:
        add_bullets(slide, 0.5, 1.1, 3.5, 4.2, "System components", ["Causal and task specification", "Evidence and model generation", "Interactive decision support", "Outcome simulation and comparison"])
        add_image(slide, "ICADOS_UCD.png", 4.35, 1.2, 6.6, 3.6)
    elif index == 19:
        add_bullets(slide, 0.5, 1.1, 3.9, 4.2, "Takeaways", ["Use tools for the decision task, not as a generic solution", "Keep causal structure visible and testable", "Support user reasoning rather than replace it", "Combine prediction, optimization, simulation, and visualization", "Evaluate recommendations for robustness, trust, and usability"])
        add_chart(slide, 5.3, 1.5, 5.6, 3.2, ["ML", "TS", "VIZ", "MO", "SIM"], [1, 3, 4, 5, 4], "Illustrative tool emphasis")

prs.save(OUT)
print(f"CREATED={OUT}")
print(f"SLIDES={len(prs.slides)}")
print(f"BYTES={OUT.stat().st_size}")
